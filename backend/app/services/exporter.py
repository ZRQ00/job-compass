import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import Job

COLUMNS = [
    ("Company", 22, lambda j: j.company),
    ("Role", 32, lambda j: j.role),
    ("Location", 20, lambda j: j.location),
    ("Status", 12, lambda j: j.status.title()),
    ("Date Applied", 14, lambda j: j.date_applied.isoformat() if j.date_applied else ""),
    ("Salary Min", 12, lambda j: j.salary_min),
    ("Salary Max", 12, lambda j: j.salary_max),
    ("Match Score", 12, lambda j: j.match_score),
    ("Notes", 40, lambda j: j.notes),
    ("Link", 40, lambda j: j.job_url),
    ("Job Description", 60, lambda j: (j.job_description or "")[:32000]),
]

STATUS_FILL = {
    "applied": "FFF4D6",
    "interview": "DDF2E3",
    "offer": "C6EBD0",
    "rejected": "F9DEDE",
    "withdrawn": "ECECEC",
}


def to_csv(jobs: list[Job]) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([c[0] for c in COLUMNS])
    for j in jobs:
        w.writerow([c[2](j) for c in COLUMNS])
    return out.getvalue()


def to_xlsx(jobs: list[Job]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (name, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
        cell.fill = PatternFill("solid", start_color="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, j in enumerate(jobs, 2):
        fill = PatternFill("solid", start_color=STATUS_FILL.get(j.status, "FFFFFF"))
        for col, (_, _, get) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r, column=col, value=get(j))
            cell.font = Font(name="Arial", size=10)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col != len(COLUMNS))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
